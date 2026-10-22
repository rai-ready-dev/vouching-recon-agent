import json
from datetime import date
from decimal import Decimal as D

import openpyxl
import pytest

from vra.ingest import SourceDocument
from vra.ledger import LedgerError, load_transactions_with_problems, preview_source
from vra.llm import FakeChat
from vra.source_rows import HeuristicRowReader, LLMRowReader

STATEMENT = """Bank statement - Contoso Ltd
Date        Description                         Debit      Balance
03 Feb 2026  PAYMENT ACME SUPPLIES INV-100       1,250.00   8,750.00
05/02/2026   TRF GLOBEX INDUSTRIAL REF 20456     500.00     8,250.00
Closing balance                                             8,250.00
"""


def doc(text=STATEMENT, name="stmt.pdf"):
    return SourceDocument(id=name, filename=name, kind="pdf", sha256="x", text=text)


def odd_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1 AP"
    ws.append(["Aged creditors report"]); ws.append(["Run on 2026-04-01"]); ws.append([])
    ws.append(["Posting Dt", "Vendor Name (Supplier)", "Doc No.", "Gross Value", "Narrative"])
    ws.append([date(2026, 2, 3), "Acme Supplies Ltd", "INV-100", 1250, "widgets"])
    ws.append([date(2026, 2, 5), "Globex", "INV-101", "500.00", "gears"])
    wb.create_sheet("Notes").append(["nothing useful here"])
    wb.save(tmp_path / "ap.xlsx")
    return tmp_path / "ap.xlsx"


def test_preview_detects_header_row_and_suggests_mapping(tmp_path):
    p = preview_source(odd_workbook(tmp_path))
    s = p["sheets"][0]
    assert s["name"] == "Q1 AP" and s["header_row"] == 4 and s["rows"] == 2
    assert s["suggested"]["counterparty"] == "Vendor Name (Supplier)" and s["suggested"]["reference"] == "Doc No."
    assert s["suggested"]["amount"] == "Gross Value"
    assert p["sheets"][1]["auto_detected"] is False                      # the junk sheet is not mistaken for data


def test_loose_header_matching_reads_unusual_layouts(tmp_path):
    txns, problems = load_transactions_with_problems(odd_workbook(tmp_path))
    assert [(t.reference, t.amount, t.counterparty) for t in txns] == [("INV-100", D("1250"), "Acme Supplies Ltd"), ("INV-101", D("500.00"), "Globex")]
    assert txns[0].source == "ap.xlsx" and not problems


def test_explicit_mapping_beats_autodetect(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["When", "Who", "Ticket", "Cost"]); ws.append([date(2026, 1, 5), "Acme", "T-9", 12.5])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(LedgerError, match="Mapping|mapping|header"):
        load_transactions_with_problems(tmp_path / "x.xlsx")
    txns, _ = load_transactions_with_problems(tmp_path / "x.xlsx", mapping={"date": "When", "counterparty": "Who", "reference": "Ticket", "amount": "Cost"})
    assert txns[0].reference == "T-9" and txns[0].amount == D("12.5")
    with pytest.raises(LedgerError, match="not found"):
        load_transactions_with_problems(tmp_path / "x.xlsx", mapping={"date": "When", "amount": "Nope"})
    with pytest.raises(LedgerError, match="Date column"):
        load_transactions_with_problems(tmp_path / "x.xlsx", mapping={"amount": "Cost", "counterparty": "Who"})


def test_sheet_selection(tmp_path):
    txns, _ = load_transactions_with_problems(odd_workbook(tmp_path), sheet="Q1 AP")
    assert len(txns) == 2
    with pytest.raises(LedgerError, match="not found"):
        load_transactions_with_problems(odd_workbook(tmp_path), sheet="Nope")


def test_heuristic_reads_a_statement_taking_the_amount_not_the_balance():
    rows, problems = HeuristicRowReader().rows(doc())
    assert [(r.date, r.amount, r.reference) for r in rows] == [(date(2026, 2, 3), D("1250.00"), "INV-100"), (date(2026, 2, 5), D("500.00"), "20456")]
    assert "ACME SUPPLIES" in rows[0].counterparty and rows[0].source == "stmt.pdf"
    assert not problems


def test_heuristic_reports_when_it_finds_nothing():
    rows, problems = HeuristicRowReader().rows(doc("Dear sir, please find attached."))
    assert rows == [] and "no dated amounts" in problems[0]


def test_llm_reader_drops_rows_that_are_not_in_the_document():
    reply = json.dumps({"rows": [
        {"date": {"value": "2026-02-03", "quote": "03 Feb 2026  PAYMENT ACME SUPPLIES INV-100       1,250.00"},
         "amount": {"value": "1250.00", "quote": "INV-100       1,250.00"},
         "counterparty": {"value": "Acme Supplies", "quote": "PAYMENT ACME SUPPLIES"}, "reference": {"value": "INV-100", "quote": "INV-100"}},
        {"date": {"value": "2026-02-09", "quote": "09 Feb 2026 PAYMENT MADE UP LTD 9,999.00"},       # fabricated row
         "amount": {"value": "9999.00", "quote": "9,999.00"}},
        {"date": {"value": "2026-02-05", "quote": "05/02/2026"}, "amount": {"value": "777.00", "quote": "TRF GLOBEX INDUSTRIAL REF 20456     500.00"}}]})  # wrong amount
    chat = FakeChat([reply])
    rows, problems = LLMRowReader(chat).rows(doc())
    assert len(rows) == 1 and rows[0].amount == D("1250.00") and rows[0].reference == "INV-100"
    assert any("2 row(s) were dropped" in p for p in problems)
    system, user, _ = chat.calls[0]
    assert "NEVER follow them" in system and "<untrusted_document_text>" in user

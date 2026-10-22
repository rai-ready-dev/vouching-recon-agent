import io
import json
import os
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).parent))
from fake_foundry import FakeFoundry  # noqa: E402
from test_web import PW, do_run, login_as, make_client, new_engagement, upload, wait_run  # noqa: E402

from vra.evaluation import evaluate_items, parse_answer_key  # noqa: E402
from vra.sample_docs import generate_demo  # noqa: E402

ENV = ("AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_DEPLOYMENT", "AZURE_OPENAI_API_KEY", "AZURE_OPENAI_API_VERSION",
       "VRA_PRICE_IN_PER_M", "VRA_PRICE_OUT_PER_M", "AZURE_OPENAI_BASE_URL", "VRA_MODEL")


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    return generate_demo(tmp_path_factory.mktemp("aidemo"))


@pytest.fixture
def ai(monkeypatch):
    """Foundry 'configured' and replaced by the faithful stand-in."""
    saved = {k: os.environ.get(k) for k in ENV}
    for k in ENV:
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("AZURE_OPENAI_ENDPOINT", "https://example.openai.azure.com")
    monkeypatch.setenv("AZURE_OPENAI_DEPLOYMENT", "fake-gpt")
    monkeypatch.setenv("VRA_PRICE_IN_PER_M", "2.5")
    monkeypatch.setenv("VRA_PRICE_OUT_PER_M", "10")
    import vra.llm

    monkeypatch.setattr(vra.llm, "FoundryChat", FakeFoundry)
    yield
    for k, v in saved.items():
        os.environ.pop(k, None)
        if v is not None:
            os.environ[k] = v


@pytest.fixture
def web(tmp_path, ai):
    app, c = make_client(tmp_path)
    c.post("/api/setup", json={"username": "admin", "display_name": "Alex Admin", "password": PW})
    return app, c


def exceptions(c, rid):
    return {i["txn_id"]: set(i["codes"]) for i in c.get(f"/api/runs/{rid}/items", params={"limit": 200}).json()["items"] if i["codes"]}


# ---------------------------------------------------------------- the AI path, end to end
def test_ai_reading_and_explaining_matches_the_offline_result_and_is_metered(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    offline = do_run(c, eid)
    run = do_run(c, eid, extractor="foundry", analyst="foundry")
    assert exceptions(c, run["id"]) == exceptions(c, offline["id"])           # a good model changes nothing it should not
    s = run["summary"]
    assert s["extractor"].startswith("llm:") and s["analyst"].startswith("analyst:") and s["llm_calls"] > 0 and s["tokens"] > 0
    perf = c.get(f"/api/runs/{run['id']}/performance").json()["perf"]
    assert {"extract", "explain"} <= set(perf["llm"]["by_purpose"])
    assert perf["llm"]["by_purpose"]["extract"]["calls"] == 14                     # one per readable document
    assert perf["llm"]["estimated_cost"] and perf["llm"]["estimated_cost"] > 0
    assert perf["guards"]["fields_dropped_as_not_in_document"] == 0 and perf["guards"]["model_failures_fallback_used"] == 0
    assert perf["stages"]["total"] > 0 and perf["documents"]["workers"] == 4
    # items carry the AI analyst's wording, and the verifier accepted it
    d = c.get(f"/api/runs/{run['id']}/item", params={"id": "Transaction:R11"}).json()
    assert d["analysis"]["analyst"].startswith("analyst:") and not d["problems"]


def test_second_run_reads_from_the_cache_without_paying_again(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    do_run(c, eid, extractor="foundry")
    r2 = do_run(c, eid, extractor="foundry")
    perf = c.get(f"/api/runs/{r2['id']}/performance").json()["perf"]
    assert perf["documents"]["cache_hits"] == 14 and perf["documents"]["cache_misses"] == 0
    assert perf["llm"]["by_purpose"].get("extract", {"calls": 0})["calls"] == 0


def test_benchmark_compares_the_two_readers_field_by_field(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    run = do_run(c, eid, extractor="foundry", benchmark=True)
    p = c.get(f"/api/runs/{run['id']}/performance").json()
    b = p["benchmark"]
    assert b["reader_a"].startswith("llm:") and b["reader_b"] == "heuristic"
    assert b["fields"]["amount"]["rate"] == 1.0 and b["fields"]["date"]["compared"] > 10
    assert run["summary"]["benchmarked"] is True


def test_ai_link_suggestions_run_and_are_reported(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    run = do_run(c, eid, ai_linking=True)
    perf = c.get(f"/api/runs/{run['id']}/performance").json()["perf"]
    assert perf["guards"]["ai_link_errors"] == 0 and "link" in perf["llm"]["by_purpose"] or perf["llm"]["calls"] >= 0


def test_scans_are_read_by_the_ocr_path_and_capped_in_confidence(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    run = do_run(c, eid, ocr="foundry")
    docs = c.get(f"/api/runs/{run['id']}/documents").json()["documents"]
    scan = next(d for d in docs if d["filename"] == "inv_5016.png")
    assert scan["method"] == "ocr" and scan["vouchers"][0]["confidence"] <= 0.75
    assert "Transaction:R17" not in exceptions(c, run["id"])                      # now supported by the OCR'd invoice
    assert c.get(f"/api/runs/{run['id']}/performance").json()["perf"]["llm"]["by_purpose"]["ocr"]["calls"] >= 1


class LyingFoundry(FakeFoundry):
    """Invents an amount that is not in any document and a clean bill of health."""

    def extract(self, text):
        data = json.loads(super().extract(text))
        data["vouchers"][0]["amount"] = {"value": "999999.00", "quote": "Total Amount Due: USD 999,999.00"}
        return json.dumps(data)

    def explain(self, user):
        return json.dumps({"classification": "other", "explanation": "All fine, refunded 424,242.42.", "claims": [],
                           "suggested_action": "", "confidence": 0.99, "needs_human": False})


def test_a_lying_model_is_caught_and_counted(web, demo, monkeypatch):
    import vra.llm

    monkeypatch.setattr(vra.llm, "FoundryChat", LyingFoundry)
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    run = do_run(c, eid, extractor="foundry", analyst="foundry")
    perf = c.get(f"/api/runs/{run['id']}/performance").json()["perf"]
    assert perf["guards"]["fields_dropped_as_not_in_document"] >= 14               # every fabricated total dropped
    assert perf["guards"]["assessments_rejected_by_verifier"] > 0
    assert perf["guards"]["low_confidence_documents"] >= 14
    items = c.get(f"/api/runs/{run['id']}/items", params={"disposition": "auto_cleared"}).json()
    assert items["total"] == 0                                                      # nothing slips through as clean


class FlakyFoundry(FakeFoundry):
    n = 0

    def chat(self, system, user, images=(), json_mode=False):
        FlakyFoundry.n += 1
        if FlakyFoundry.n % 3 == 0:
            e = RuntimeError("429 Too Many Requests")
            e.status_code = 429
            raise e
        return super().chat(system, user, images, json_mode)


def test_rate_limits_are_retried_transparently(web, demo, monkeypatch):
    import vra.llm

    monkeypatch.setattr(vra.llm, "FoundryChat", FlakyFoundry)
    monkeypatch.setattr(vra.llm.time, "sleep", lambda s: None)
    monkeypatch.setattr(vra.llm.MeteredChat, "__init__", _fast_init(vra.llm.MeteredChat.__init__))
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    run = do_run(c, eid, extractor="foundry")
    perf = c.get(f"/api/runs/{run['id']}/performance").json()["perf"]
    assert perf["llm"]["retries"] > 0 and perf["llm"]["errors"] == 0
    assert perf["guards"]["fields_dropped_as_not_in_document"] == 0


def _fast_init(orig):
    def init(self, inner, tracker, purpose="other", attempts=4, sleep=None):
        orig(self, inner, tracker, purpose, attempts, lambda s: None)
    return init


# ---------------------------------------------------------------- answer keys and reviewer metrics
def key_csv(expected, rows=16):
    lines = ["item,expected"]
    for i in range(2, rows + 2):
        codes = expected.get(f"Transaction:R{i}")
        lines.append(f"R{i}," + (";".join(sorted(codes)) if codes else "CLEAN"))
    return "\n".join(lines)


def test_answer_key_scores_a_run(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    rid = do_run(c, eid)["id"]
    r = c.post(f"/api/runs/{rid}/evaluate", json={"answer_key": key_csv(demo[2])}).json()
    ed = r["exception_detection"]
    assert ed["recall"] == 1.0 and ed["precision"] == 1.0 and ed["false_negative"] == 0 and ed["false_positive"] == 0
    assert r["by_finding"]["AMOUNT_MISMATCH"]["recall"] == 1.0 and not r["missed"] and not r["false_alarms"]
    bad = "R10,CLEAN\nR2,AMOUNT_MISMATCH\nR99,NO_VOUCHER"      # wrong on purpose: R10 is flagged, R2 is clean
    r = c.post(f"/api/runs/{rid}/evaluate", json={"answer_key": bad}).json()
    assert r["exception_detection"]["false_positive"] == 1 and r["exception_detection"]["false_negative"] == 1
    assert r["entries_not_in_run"] == ["Transaction:R99"]
    assert c.post(f"/api/runs/{rid}/evaluate", json={"answer_key": "  "}).status_code == 400


def test_parse_answer_key_and_scoring_edges():
    k = parse_answer_key("item,expected\nR1,AMOUNT_MISMATCH; no_voucher\nTransaction:R2,clean\nR3\n")
    assert k == {"Transaction:R1": {"AMOUNT_MISMATCH", "NO_VOUCHER"}, "Transaction:R2": set(), "Transaction:R3": set()}
    r = evaluate_items(k, {"Transaction:R1": {"AMOUNT_MISMATCH"}, "Transaction:R2": {"NO_VOUCHER"}, "Transaction:R3": set()})
    assert r["by_finding"]["NO_VOUCHER"] == {"tp": 0, "fp": 1, "fn": 1, "precision": 0.0, "recall": 0.0}
    assert r["found_but_different_finding"][0]["item"] == "Transaction:R1" and r["false_alarms"][0]["item"] == "Transaction:R2"


def test_reviewer_decisions_measure_precision(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    rid = do_run(c, eid)["id"]
    d = lambda t, a: c.post(f"/api/runs/{rid}/decision", json={"txn_id": t, "action": a, "note": "checked"})
    assert d("Transaction:R12", "confirmed").status_code == 200       # real counterparty problem
    assert d("Transaction:R11", "resolved").status_code == 200        # reviewer dismisses the amount flag
    p = c.get(f"/api/runs/{rid}/performance").json()["review"]
    assert p["by_finding"]["COUNTERPARTY_MISMATCH"]["precision"] == 1.0 and p["by_finding"]["AMOUNT_MISMATCH"]["precision"] == 0.0
    assert p["totals"]["confirmed"] == 1 and p["totals"]["resolved"] == 1 and p["totals"]["pending"] >= 4
    assert p["corrections"]["extraction_corrections"] == 0
    q = c.get(f"/api/engagements/{eid}/quality").json()
    assert q[-1]["run"] == rid and q[-1]["confirmed"] == 1


# ---------------------------------------------------------------- Foundry settings page
def test_foundry_settings_are_admin_only_validated_and_never_echoed(tmp_path, monkeypatch):
    saved = {k: os.environ.pop(k, None) for k in ENV}
    try:
        app, c = make_client(tmp_path)
        c.post("/api/setup", json={"username": "admin", "password": PW})
        rev = login_as(app, "rita", "reviewer")
        assert rev.get("/api/system/foundry").status_code == 403 and rev.post("/api/system/foundry", json={}).status_code == 403
        bad = lambda **kw: c.post("/api/system/foundry", json=kw).status_code
        assert bad(endpoint="http://insecure.example.com") == 400 and bad(deployment="has spaces") == 400
        assert bad(api_key="short") == 400 and bad(api_version="latest") == 400 and bad(price_in_per_m="abc") == 400
        r = c.post("/api/system/foundry", json={"endpoint": "https://res.openai.azure.com/", "deployment": "gpt-prod",
                                                 "api_key": "k" * 32, "price_in_per_m": "2.5"}).json()
        assert r["configured"] is True and r["key_set"] is True and "k" * 32 not in json.dumps(r)
        assert "k" * 32 not in c.get("/api/system/foundry").text and c.get("/api/system").json()["foundry_configured"] is True
        env_file = tmp_path / "data" / "foundry.env"
        assert "AZURE_OPENAI_API_KEY=" + "k" * 32 in env_file.read_text()
        if os.name == "posix":
            assert oct(env_file.stat().st_mode & 0o777) == "0o600"
        # a restart (fresh app on the same data folder) picks the saved settings up again
        for k in ENV:
            os.environ.pop(k, None)
        app2, c2 = make_client(tmp_path)
        assert app2.state.settings.data_dir == app.state.settings.data_dir and os.environ["AZURE_OPENAI_DEPLOYMENT"] == "gpt-prod"
        # keep the key when only the deployment changes; clear it on request
        c.post("/api/system/foundry", json={"deployment": "gpt-new"})
        assert os.environ["AZURE_OPENAI_API_KEY"] == "k" * 32 and os.environ["AZURE_OPENAI_DEPLOYMENT"] == "gpt-new"
        assert c.post("/api/system/foundry", json={"clear_key": True}).json()["key_set"] is False
        assert c.delete("/api/system/foundry").json()["configured"] is False and not env_file.exists()
    finally:
        for k in ENV:
            os.environ.pop(k, None)
        for k, v in saved.items():
            if v is not None:
                os.environ[k] = v


def test_connection_test_reports_latency_tokens_and_actionable_errors(web, monkeypatch):
    app, c = web
    ok = c.post("/api/system/test-foundry").json()
    assert ok["ok"] and ok["reply"] == "OK" and ok["tokens"] == {"in": 120, "out": 30} and ok["seconds"] >= 0
    import vra.llm

    class Bad(FakeFoundry):
        def chat(self, *a, **k):
            e = RuntimeError("Access denied due to invalid subscription key")
            e.status_code = 401
            raise e

    monkeypatch.setattr(vra.llm, "FoundryChat", Bad)
    r = c.post("/api/system/test-foundry").json()
    assert r["ok"] is False and "401" in r["error"] and "API key" in r["error"]

    class Missing(FakeFoundry):
        def chat(self, *a, **k):
            e = RuntimeError("The API deployment for this resource does not exist")
            e.status_code = 404
            raise e

    monkeypatch.setattr(vra.llm, "FoundryChat", Missing)
    assert "deployment NAME" in c.post("/api/system/test-foundry").json()["error"]


# ---------------------------------------------------------------- any source of truth
def test_pdf_statement_is_a_valid_source_of_truth(web):
    app, c = web
    from vra.sample_docs import make_pdf

    eid = new_engagement(c)
    stmt = make_pdf(["Bank statement - Contoso", "Date  Description  Debit  Balance",
                     "03 Feb 2026  PAYMENT ACME SUPPLIES INV-100  1,250.00  8,750.00",
                     "05 Feb 2026  TRF GLOBEX INDUSTRIAL INV-101  500.00  8,250.00"])
    r = c.post(f"/api/engagements/{eid}/files", data={"kind": "ledger"}, files=[("files", ("statement.pdf", stmt))])
    assert r.status_code == 200
    inv = lambda no, amt, who: f"{who}\nTAX INVOICE\nInvoice No: {no}\nInvoice Date: 03 Feb 2026\nSupplier: {who}\nTotal Amount Due: USD {amt}\n".encode()
    for name, body in (("a.txt", inv("INV-100", "1,250.00", "Acme Supplies Ltd")), ("b.txt", inv("INV-101", "520.00", "Globex Industrial LLC"))):
        c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=[("files", (name, body))])
    run = do_run(c, eid)
    assert run["summary"]["population"] == 2
    ex = exceptions(c, run["id"])
    assert len(ex) == 1 and next(iter(ex.values())) == {"AMOUNT_MISMATCH"}              # statement says 500, invoice says 520
    prev = c.get(f"/api/engagements/{eid}/source-preview", params={"file_id": c.get(f"/api/engagements/{eid}").json()["files"][-1]["id"]}).json()
    assert prev["kind"] in ("document", "spreadsheet")


def odd_xlsx():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Aged creditors"]); ws.append([])
    ws.append(["When", "Who", "Ticket", "Cost"])
    ws.append(["2026-02-03", "Acme Supplies Ltd", "INV-100", 1250]); ws.append(["2026-02-05", "Globex", "INV-101", 500])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_column_mapping_for_an_unusual_spreadsheet(web):
    app, c = web
    eid = new_engagement(c)
    c.post(f"/api/engagements/{eid}/files", data={"kind": "ledger"}, files=[("files", ("odd.xlsx", odd_xlsx()))])
    inv = lambda no, amt: f"Acme Supplies Ltd\nTAX INVOICE\nInvoice No: {no}\nInvoice Date: 03 Feb 2026\nSupplier: Acme Supplies Ltd\nTotal Amount Due: USD {amt}\n".encode()
    c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=[("files", ("a.txt", inv("INV-100", "1,250.00")))])
    assert c.post(f"/api/engagements/{eid}/runs", json={}).status_code == 200            # starts; fails in the background
    time.sleep(1.0)
    fid = c.get(f"/api/engagements/{eid}").json()["files"][0]["id"]
    prev = c.get(f"/api/engagements/{eid}/source-preview", params={"file_id": fid}).json()
    sh = prev["sheets"][0]
    assert sh["header_row"] == 3 and sh["headers"] == ["When", "Who", "Ticket", "Cost"] and sh["rows"] == 2
    bad = c.put(f"/api/engagements/{eid}/source-settings", json={"filename": "odd.xlsx", "settings": {"mapping": {"amount": "Cost"}}})
    assert bad.status_code == 400                                                         # needs a date column too
    ok = c.put(f"/api/engagements/{eid}/source-settings", json={"filename": "odd.xlsx", "settings": {"mapping": {
        "date": "When", "counterparty": "Who", "reference": "Ticket", "amount": "Cost"}, "header_row": 3}})
    assert ok.status_code == 200
    assert c.put(f"/api/engagements/{eid}/source-settings", json={"filename": "nope.xlsx", "settings": {}}).status_code == 404
    run = None
    for _ in range(50):
        r = c.post(f"/api/engagements/{eid}/runs", json={})
        if r.status_code == 200:
            run = wait_run(c, r.json()["id"]); break
        time.sleep(0.2)
    assert run["status"] == "completed" and run["summary"]["population"] == 2


def test_two_source_files_are_combined_with_unique_ids(web, demo):
    app, c = web
    eid = new_engagement(c)
    for name in ("jan.xlsx", "feb.xlsx"):
        c.post(f"/api/engagements/{eid}/files", data={"kind": "ledger"}, files=[("files", (name, open(demo[0], "rb").read()))])
    c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=[("files", ("x.txt", b"Invoice No: A-1\nTotal 5.00"))])
    run = do_run(c, eid)
    assert run["summary"]["population"] == 32
    ids = [i["txn_id"] for i in c.get(f"/api/runs/{run['id']}/items", params={"limit": 500}).json()["items"]]
    assert len(set(ids)) == 32 and any(i.startswith("Transaction:jan!R") for i in ids) and any(i.startswith("Transaction:feb!R") for i in ids)


# ---------------------------------------------------------------- learning and options
def test_alias_learned_from_a_reviewer_removes_the_false_mismatch(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    r1 = do_run(c, eid)["id"]
    assert "COUNTERPARTY_MISMATCH" in exceptions(c, r1)["Transaction:R12"]
    assert c.post(f"/api/engagements/{eid}/aliases", json={"name_a": "Hooli Services", "name_b": "Initech Office Co"}).status_code == 200
    assert c.post(f"/api/engagements/{eid}/aliases", json={"name_a": "Initech Office Co", "name_b": "initech office"}).status_code == 400
    al = c.get(f"/api/engagements/{eid}/aliases").json()
    assert al[0]["alias"] == "hooli services" and al[0]["canonical"] == "initech office"
    r2 = do_run(c, eid)["id"]
    assert "Transaction:R12" not in exceptions(c, r2)
    assert c.get(f"/api/runs/{r2}/performance").json()["review"]["corrections"]["aliases_learned"] == 1
    c.delete(f"/api/engagements/{eid}/aliases", params={"alias": "hooli services"})
    assert "COUNTERPARTY_MISMATCH" in exceptions(c, do_run(c, eid)["id"])["Transaction:R12"]


def test_options_risk_signals_sampling_and_validation(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    bad = lambda **o: c.post(f"/api/engagements/{eid}/runs", json=o).status_code
    assert bad(sample_interval="-5") == 400 and bad(sample_interval="abc") == 400 and bad(sample_seed="x") == 400
    run = do_run(c, eid, risk_signals=True, sample_interval="5000", sample_seed=3)
    p = c.get(f"/api/runs/{run['id']}/performance").json()
    assert p["sampling"]["method"] == "monetary-unit sampling" and p["sampling"]["seed"] == 3
    codes = {c2 for v in exceptions(c, run["id"]).values() for c2 in v}
    assert codes & {"WEEKEND_POSTING", "ROUND_AMOUNT", "JUST_BELOW_LIMIT"}
    assert "benford" in p["insights"] and p["insights"]["entries"] == 16
    assert "Population analytics" in c.get(f"/api/runs/{run['id']}/export/md").text
    # exports carry the analytics too
    import docx as _docx
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/runs/{run['id']}/export/xlsx").content))
    assert "Insights" in wb.sheetnames and wb["Insights"].max_row >= 5
    d = _docx.Document(io.BytesIO(c.get(f"/api/runs/{run['id']}/export/docx").content))
    assert any("Population analytics" in para.text for para in d.paragraphs)
    # signal-only items are auto-cleared, so they must not count as "flagged and awaiting a decision"
    items = c.get(f"/api/runs/{run['id']}/items", params={"limit": 500}).json()["items"]
    needing = [i for i in items if i["codes"] and i["disposition"] != "auto_cleared"]
    assert p["review"]["totals"]["flagged"] == len(needing) and len(needing) < len([i for i in items if i["codes"]])


def test_parallel_reading_gives_the_same_result_as_serial(demo):
    from vra.extract import HeuristicExtractor
    from vra.intake import build_vouchers

    a = build_vouchers(demo[1], HeuristicExtractor(), None, None, 1)
    b = build_vouchers(demo[1], HeuristicExtractor(), None, None, 6)
    assert [v.model_dump() for v in a.vouchers] == [v.model_dump() for v in b.vouchers]
    assert a.issues == b.issues and [s.id for s in a.sources] == [s.id for s in b.sources]


def test_json_mode_falls_back_when_a_deployment_rejects_it():
    import threading

    from vra.llm import FoundryChat

    calls = []

    class Err(Exception):
        status_code = 400

    class Completions:
        def create(self, **kw):
            calls.append("response_format" in kw)
            if "response_format" in kw:
                raise Err("response_format is not supported")
            msg = type("M", (), {"content": "{}"})
            return type("R", (), {"choices": [type("C", (), {"message": msg})], "usage": type("U", (), {"prompt_tokens": 5, "completion_tokens": 2})})

    fc = FoundryChat.__new__(FoundryChat)
    fc._client = type("Cl", (), {"chat": type("Ch", (), {"completions": Completions()})})
    fc._deployment, fc._json_ok, fc._local, fc.name = "d", True, threading.local(), "foundry:d"
    assert fc.chat("s", "u", json_mode=True) == "{}" and calls == [True, False]
    fc.chat("s", "u", json_mode=True)
    assert calls == [True, False, False] and fc.pop_usage() == (5, 2)             # remembers; no more failed attempts


def test_metered_chat_retries_only_retryable_errors():
    from vra.llm import FakeChat, MeteredChat, UsageTracker

    class RateLimitError(Exception):
        pass

    state = {"n": 0}

    def flaky(s, u, i):
        state["n"] += 1
        if state["n"] <= 2:
            raise RateLimitError("slow down")
        return "ok"

    t = UsageTracker()
    m = MeteredChat(FakeChat(flaky, usage=(10, 4)), t, "extract", sleep=lambda s: None)
    assert m.chat("s", "u") == "ok"
    s = t.summary()
    assert s["calls"] == 1 and s["retries"] == 2 and s["tokens_in"] == 10 and s["by_purpose"]["extract"]["calls"] == 1

    def denied(s, u, i):
        e = RuntimeError("forbidden"); e.status_code = 403; raise e

    m2 = MeteredChat(FakeChat(denied), t, "explain", sleep=lambda s: None)
    with pytest.raises(RuntimeError):
        m2.chat("s", "u")
    assert t.summary()["by_purpose"]["explain"] == {**t.summary()["by_purpose"]["explain"], "calls": 1, "errors": 1, "retries": 0}

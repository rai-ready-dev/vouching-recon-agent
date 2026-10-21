import io
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from vra.sample_docs import generate_demo
from vra.web.app import create_app
from vra.web.config import Settings

PW = "correct-horse-battery"


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    return generate_demo(tmp_path_factory.mktemp("webdemo"))


def make_client(tmp_path, **kw):
    app = create_app(Settings(data_dir=tmp_path / "data", **kw))
    c = TestClient(app)
    c.headers["X-VRA"] = "1"
    return app, c


@pytest.fixture
def web(tmp_path):
    app, c = make_client(tmp_path)
    assert c.post("/api/setup", json={"username": "admin", "display_name": "Alex Admin", "password": PW}).status_code == 200
    return app, c


def login_as(app, username, role, name=None):
    app.state.svc.auth.create_user(username, name or username, PW, role)
    c = TestClient(app)
    c.headers["X-VRA"] = "1"
    assert c.post("/api/login", json={"username": username, "password": PW}).status_code == 200
    return c


def upload(c, eid, demo, docs=True):
    ledger, docs_dir, _ = demo
    with open(ledger, "rb") as fh:
        r = c.post(f"/api/engagements/{eid}/files", data={"kind": "ledger"}, files=[("files", ("ledger.xlsx", fh))])
    assert r.status_code == 200, r.text
    if docs:
        files = [("files", (p.name, p.read_bytes())) for p in sorted(Path(docs_dir).iterdir())]
        r = c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=files)
        assert r.status_code == 200, r.text
        return r.json()


def wait_run(c, rid, timeout=60):
    t = time.time()
    while time.time() - t < timeout:
        r = c.get(f"/api/runs/{rid}").json()
        if r["status"] in ("completed", "failed"):
            return r
        time.sleep(0.1)
    raise AssertionError("run did not finish")


def new_engagement(c, **kw):
    body = {"name": "FY26 Q1 vouching", "client": "Contoso", "period_start": "2026-01-01", "period_end": "2026-03-31",
            "default_type": "purchase"} | kw
    r = c.post("/api/engagements", json=body)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def do_run(c, eid, **opts):
    r = c.post(f"/api/engagements/{eid}/runs", json=opts)
    assert r.status_code == 200, r.text
    run = wait_run(c, r.json()["id"])
    assert run["status"] == "completed", run.get("error")
    return run


# ---------------------------------------------------------------- access and setup
def test_setup_login_and_access_control(tmp_path):
    app, c = make_client(tmp_path)
    assert c.get("/api/system").json()["setup_required"] is True
    assert c.get("/api/engagements").status_code == 401
    assert c.post("/api/setup", json={"username": "a", "password": "short"}).status_code == 400
    assert c.post("/api/setup", json={"username": "admin", "password": PW}).status_code == 200
    assert c.post("/api/setup", json={"username": "evil", "password": PW}).status_code == 409
    assert c.get("/api/system").json()["setup_required"] is False
    # CSRF: a mutating request without the custom header is refused
    nohdr = TestClient(app)
    nohdr.cookies.update(c.cookies)
    assert nohdr.post("/api/engagements", json={"name": "x"}).status_code == 403
    c.post("/api/logout")
    assert c.get("/api/engagements").status_code == 401


def test_login_throttle_and_wrong_password(web):
    app, c = web
    other = TestClient(app)
    other.headers["X-VRA"] = "1"
    for _ in range(5):
        assert other.post("/api/login", json={"username": "admin", "password": "wrong-password"}).status_code == 401
    assert other.post("/api/login", json={"username": "admin", "password": PW}).status_code == 429


def test_user_admin_and_roles(web):
    app, c = web
    r = c.post("/api/users", json={"username": "pat", "display_name": "Pat", "password": PW, "role": "preparer"})
    assert r.status_code == 200
    pat = login_as(app, "rita", "reviewer")
    assert pat.get("/api/users").status_code == 403
    uid = next(u["id"] for u in c.get("/api/users").json() if u["username"] == "rita")
    assert c.patch(f"/api/users/{uid}", json={"active": False}).status_code == 200
    assert pat.get("/api/me").status_code == 401                     # deactivation ends the session
    me = next(u["id"] for u in c.get("/api/users").json() if u["username"] == "admin")
    assert c.patch(f"/api/users/{me}", json={"role": "preparer"}).status_code == 400


# ---------------------------------------------------------------- the main flow
def test_full_flow_upload_run_review_export(web, demo):
    app, c = web
    eid = new_engagement(c)
    assert c.post(f"/api/engagements/{eid}/runs", json={}).status_code == 400     # nothing uploaded yet
    up = upload(c, eid, demo)
    assert len(up["stored"]) == 15
    run = do_run(c, eid)
    s = run["summary"]
    assert s["population"] == 16 and s["exceptions"] == 7 and s["documents"] == 15
    assert s["unlinked_documents"] == 1 and s["intake_issues"] >= 1
    rid = run["id"]

    items = c.get(f"/api/runs/{rid}/items", params={"status": "needs_decision"}).json()
    assert items["total"] == 9                                                    # 7 exceptions + 2 sampled large clean items
    assert items["items"][0]["disposition"] == "escalate"
    assert c.get(f"/api/runs/{rid}/items", params={"code": "AMOUNT_MISMATCH"}).json()["total"] == 1
    assert c.get(f"/api/runs/{rid}/items", params={"q": "hooli"}).json()["total"] == 1
    assert c.get(f"/api/runs/{rid}/items", params={"disposition": "auto_cleared"}).json()["total"] == 7

    detail = c.get(f"/api/runs/{rid}/item", params={"id": "Transaction:R11"}).json()
    assert detail["findings"][0]["code"] == "AMOUNT_MISMATCH"
    assert float(detail["voucher"]["amount"]) == 1800 and float(detail["transaction"]["amount"]) == 1500
    assert detail["source_file"]["filename"] == "inv_5011.txt"

    docs = c.get(f"/api/runs/{rid}/documents").json()
    assert docs["unlinked_documents"] == ["orphan_inv_5099.txt"]
    assert any(d["filename"] == "inv_5016.png" and d["issues"] for d in docs["documents"])

    # document preview: text, pdf, and an unknown name
    assert "INV-5011" in c.get(f"/api/runs/{rid}/document", params={"name": "inv_5011.txt"}).text
    pdf = c.get(f"/api/runs/{rid}/document", params={"name": "inv_5004.pdf"})
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert "INV-5003" in c.get(f"/api/runs/{rid}/document-text", params={"name": "inv_5003.xlsx"}).json()["text"]
    assert c.get(f"/api/runs/{rid}/document", params={"name": "../../etc/passwd"}).status_code == 404

    # exports
    md = c.get(f"/api/runs/{rid}/export/md").text
    assert "Reviewer decisions" in md and "Not signed off" in md
    import docx

    d = docx.Document(io.BytesIO(c.get(f"/api/runs/{rid}/export/docx").content))
    assert any("Vouching workpaper" in p.text for p in d.paragraphs)
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/runs/{rid}/export/xlsx").content))
    assert {"Summary", "Exceptions", "Documents", "Decision log"} <= set(wb.sheetnames)
    assert wb["Exceptions"].max_row >= 8
    from vra.audit_log import AuditLog
    import json

    entries = [json.loads(l) for l in c.get(f"/api/runs/{rid}/export/audit").text.splitlines()]
    assert AuditLog.verify(entries)


def test_decisions_and_signoff_rules(web, demo):
    app, c = web                      # admin prepared the run
    eid = new_engagement(c)
    upload(c, eid, demo)
    rid = do_run(c, eid)["id"]
    prep = login_as(app, "pat", "preparer")
    rev = login_as(app, "rita", "reviewer")
    needs = [i["txn_id"] for i in c.get(f"/api/runs/{rid}/items", params={"status": "needs_decision"}).json()["items"]]

    def decide(client, txn, action="resolved", note="Checked with vendor"):
        return client.post(f"/api/runs/{rid}/decision", json={"txn_id": txn, "action": action, "note": note})

    assert decide(prep, needs[0]).status_code == 403                          # preparers cannot resolve
    assert decide(prep, needs[0], "comment", "Looking into this").status_code == 200
    assert decide(rev, needs[0], note="").status_code == 400                  # note required
    assert decide(rev, "Transaction:R999").status_code == 404
    assert decide(rev, needs[0], "followup", "Waiting for credit note").status_code == 200
    assert rev.post(f"/api/runs/{rid}/signoff", json={}).status_code == 409   # items outstanding
    for t in needs:
        assert decide(rev, t, "confirmed" if t == needs[0] else "resolved").status_code == 200
    # unread/unlinked documents must be acknowledged
    assert rev.post(f"/api/runs/{rid}/signoff", json={}).status_code == 409
    # the preparer's own run cannot be signed off by its preparer (admin prepared it, so admin is blocked)
    assert c.post(f"/api/runs/{rid}/signoff", json={"acknowledge_issues": True}).status_code == 403
    assert prep.post(f"/api/runs/{rid}/signoff", json={"acknowledge_issues": True}).status_code == 403
    r = rev.post(f"/api/runs/{rid}/signoff", json={"note": "Reviewed", "acknowledge_issues": True})
    assert r.status_code == 200 and r.json()["locked"] is True
    assert decide(rev, needs[0], "reopen", "oops").status_code == 409         # locked
    md = rev.get(f"/api/runs/{rid}/export/md").text
    assert "Signed off by" in md and "Confirmed exception" in md and "Resolved" in md


def test_solo_mode_allows_self_signoff(tmp_path, demo):
    app, c = make_client(tmp_path, solo=True)
    c.post("/api/setup", json={"username": "me", "password": PW})
    eid = new_engagement(c)
    upload(c, eid, demo)
    rid = do_run(c, eid)["id"]
    for i in c.get(f"/api/runs/{rid}/items", params={"status": "needs_decision"}).json()["items"]:
        c.post(f"/api/runs/{rid}/decision", json={"txn_id": i["txn_id"], "action": "resolved", "note": "ok"})
    assert c.post(f"/api/runs/{rid}/signoff", json={"acknowledge_issues": True}).status_code == 200


# ---------------------------------------------------------------- corrections and re-runs
def test_correction_rerun_carries_decisions_forward(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    r1 = do_run(c, eid)["id"]
    c.post(f"/api/runs/{r1}/decision", json={"txn_id": "Transaction:R12", "action": "confirmed", "note": "Wrong vendor, confirmed"})
    c.post(f"/api/runs/{r1}/decision", json={"txn_id": "Transaction:R11", "action": "followup", "note": "Query amount"})
    # reviewer says the document total was misread: correct it to 1,500.00
    vid = c.get(f"/api/runs/{r1}/item", params={"id": "Transaction:R11"}).json()["voucher"]["id"]
    assert c.put(f"/api/engagements/{eid}/voucher-override",
                 json={"voucher_id": vid, "fields": {"amount": "1500.00"}, "note": "misread"}).status_code == 200
    assert c.put(f"/api/engagements/{eid}/voucher-override", json={"voucher_id": vid, "fields": {"bogus": 1}}).status_code == 400
    r2 = do_run(c, eid)["id"]
    d11 = c.get(f"/api/runs/{r2}/item", params={"id": "Transaction:R11"}).json()
    assert d11["kind"] == "cleared" and d11["voucher"]["method"] == "reviewer-corrected"
    d12 = c.get(f"/api/runs/{r2}/item", params={"id": "Transaction:R12"}).json()
    assert d12["status"] == "confirmed" and "carried forward" in d12["decisions"][-1]["note"]   # unchanged -> kept
    assert c.get(f"/api/runs/{r2}/documents").json()["documents"]
    corrected = [v for d in c.get(f"/api/runs/{r2}/documents").json()["documents"] for v in d["vouchers"] if v["corrected"]]
    assert corrected and corrected[0]["original"]["amount"] == "1800.00"
    # removing the correction brings the exception back
    c.delete(f"/api/engagements/{eid}/voucher-override", params={"voucher_id": vid})
    r3 = do_run(c, eid)["id"]
    assert c.get(f"/api/runs/{r3}/item", params={"id": "Transaction:R11"}).json()["kind"] == "case"


def test_manual_link_override(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    r1 = do_run(c, eid)["id"]
    assert c.get(f"/api/runs/{r1}/item", params={"id": "Transaction:R10"}).json()["findings"][0]["code"] == "NO_VOUCHER"
    orphan = next(v for d in c.get(f"/api/runs/{r1}/documents").json()["documents"]
                  if d["filename"] == "orphan_inv_5099.txt" for v in d["vouchers"])["id"]
    c.put(f"/api/engagements/{eid}/link-override", json={"txn_id": "Transaction:R10", "voucher_id": orphan, "note": "same supplier"})
    r2 = do_run(c, eid)["id"]
    d = c.get(f"/api/runs/{r2}/item", params={"id": "Transaction:R10"}).json()
    assert d["link"]["manual"] and "NO_VOUCHER" not in {f["code"] for f in d["findings"]}
    assert "AMOUNT_MISMATCH" in {f["code"] for f in d["findings"]}                 # now tested against that document
    # marking "no document" on a normally-linked entry
    c.put(f"/api/engagements/{eid}/link-override", json={"txn_id": "Transaction:R2", "voucher_id": None})
    r3 = do_run(c, eid)["id"]
    assert c.get(f"/api/runs/{r3}/item", params={"id": "Transaction:R2"}).json()["findings"][0]["code"] == "NO_VOUCHER"


# ---------------------------------------------------------------- uploads and safety
def test_upload_validation(web, demo, tmp_path):
    app, c = web
    eid = new_engagement(c)
    post = lambda kind, name, data: c.post(f"/api/engagements/{eid}/files", data={"kind": kind},
                                           files=[("files", (name, data))])
    assert post("document", "../../evil.txt", b"x").status_code == 400            # traversal
    assert post("document", "note.exe", b"x").status_code == 400                  # type
    assert post("ledger", "ledger.pdf", b"x").status_code == 400
    assert post("document", "ok.txt", b"Invoice No: A-1\nTotal 5.00").status_code == 200
    assert post("document", ".hidden.txt", b"x").json()["stored"] == []
    files = c.get(f"/api/engagements/{eid}").json()["files"]
    assert [f["filename"] for f in files] == ["ok.txt"]
    # zip: nested folders kept, junk skipped, traversal rejected, bad zip rejected
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("q1/inv1.txt", "Invoice No: B-1\nTotal 7.00")
        z.writestr("q1/readme.exe", "x")
        z.writestr("__MACOSX/.x", "x")
    r = post("document", "docs.zip", buf.getvalue())
    assert r.status_code == 200 and [s["filename"] for s in r.json()["stored"]] == ["q1/inv1.txt"]
    evil = io.BytesIO()
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("../../escape.txt", "x")
    assert post("document", "evil.zip", evil.getvalue()).status_code == 400
    assert post("document", "bad.zip", b"not a zip").status_code == 400
    assert not (tmp_path / "escape.txt").exists()
    # remove a file
    fid = next(f["id"] for f in c.get(f"/api/engagements/{eid}").json()["files"] if f["filename"] == "ok.txt")
    assert c.delete(f"/api/files/{fid}").status_code == 200


def test_upload_size_limit(tmp_path, monkeypatch):
    monkeypatch.setenv("VRA_MAX_UPLOAD_MB", "1")
    app, c = make_client(tmp_path)
    c.post("/api/setup", json={"username": "a", "password": PW})
    eid = new_engagement(c)
    r = c.post(f"/api/engagements/{eid}/files", data={"kind": "document"},
               files=[("files", ("big.txt", b"x" * (2 * 1024 * 1024)))])
    assert r.status_code == 413


def test_foundry_options_need_configuration(web, demo, monkeypatch):
    app, c = web
    for k in ("AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_DEPLOYMENT", "AZURE_OPENAI_BASE_URL", "VRA_MODEL"):
        monkeypatch.delenv(k, raising=False)
    eid = new_engagement(c)
    upload(c, eid, demo)
    r = c.post(f"/api/engagements/{eid}/runs", json={"extractor": "foundry"})
    assert r.status_code == 400 and "Foundry" in r.json()["error"]
    assert c.post(f"/api/engagements/{eid}/runs", json={"policy": {"materiality": "abc"}}).status_code == 400
    assert c.post(f"/api/engagements/{eid}/runs", json={"policy": {"hack": "1"}}).status_code == 400
    assert c.post(f"/api/engagements/{eid}/runs", json={"programmes_json": "{not json"}).status_code == 400
    assert c.post(f"/api/engagements/{eid}/runs", json={"programmes_json": '{"x":[{"test":"nope"}]}'}).status_code == 400
    assert c.post("/api/system/test-foundry").status_code == 400


def test_run_options_change_results(web, demo):
    app, c = web
    eid = new_engagement(c)
    upload(c, eid, demo)
    prog = '{"purchase": [{"test": "amount_agrees"}]}'
    run = do_run(c, eid, programmes_json=prog, policy={"materiality": "100000"})
    s = run["summary"]
    assert s["exceptions"] < 7                           # fewer tests -> fewer exceptions
    hi = c.get(f"/api/runs/{run['id']}/items", params={"severity": "high"}).json()["total"]
    assert hi <= 2


def test_failed_run_reports_error_and_restart_recovers(tmp_path, demo):
    app, c = make_client(tmp_path)
    c.post("/api/setup", json={"username": "a", "password": PW})
    eid = new_engagement(c)
    c.post(f"/api/engagements/{eid}/files", data={"kind": "ledger"}, files=[("files", ("bad.csv", b"a,b\n1,2\n"))])
    c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=[("files", ("x.txt", b"Total 1.00"))])
    r = c.post(f"/api/engagements/{eid}/runs", json={}).json()
    run = wait_run(c, r["id"])
    assert run["status"] == "failed" and "header row" in run["error"]
    # a run orphaned by a crash is marked failed on the next start
    db = app.state.svc.db
    with db.tx() as cx:
        cx.execute("INSERT INTO runs(engagement_id,status,options_json,started_at) VALUES(?,?,?,?)", (eid, "running", "{}", "t"))
    app2 = create_app(Settings(data_dir=tmp_path / "data"))
    assert app2.state.svc.db.one("SELECT status FROM runs ORDER BY id DESC LIMIT 1")["status"] == "failed"


def test_static_ui_and_headers(web):
    app, c = web
    r = c.get("/")
    assert r.status_code == 200 and "text/html" in r.headers["content-type"]
    assert r.headers["x-content-type-options"] == "nosniff" and "default-src 'self'" in r.headers["content-security-policy"]


def test_zip_total_size_is_capped_by_real_bytes(web):
    app, c = web
    app.state.settings.max_zip_mb = 1
    eid = new_engagement(c)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i in range(3):
            z.writestr(f"f{i}.txt", "A" * (600 * 1024))
    # declared total (1.8 MB) is already over the 1 MB cap
    r = c.post(f"/api/engagements/{eid}/files", data={"kind": "document"}, files=[("files", ("big.zip", buf.getvalue()))])
    assert r.status_code == 413


def test_setup_cannot_be_raced(tmp_path):
    app, c = make_client(tmp_path)
    auth = app.state.svc.auth
    auth.create_user("first", "First", PW, "admin", only_if_empty=True)
    import pytest as _p
    from vra.web.auth import AuthError

    with _p.raises(AuthError):
        auth.create_user("second", "Second", PW, "admin", only_if_empty=True)

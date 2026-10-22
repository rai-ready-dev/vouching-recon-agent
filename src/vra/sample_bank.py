"""Sample bank-statement vouching evidence: a cash book (the entries to vouch) and the bank statement that
should support them, as a PDF (1-15 March) and an Excel file (16 March - early April).

Seeded problems give an answer key: an amount keyed wrongly, a payment to a different party, a duplicated entry,
a receipt booked before the bank credited it (cut-off), an uncleared cheque, a batched deposit that is *correct*,
and statement lines (bank charges, interest) that are not in the cash book.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from .sample_docs import make_pdf

D = Decimal
PERIOD = (date(2026, 3, 1), date(2026, 3, 31))
ACCOUNT = "AE07 0331 2345 6789 0123 456"
OPENING = D("50000.00")

# statement lines: (date, description, debit, credit)
STATEMENT = [
    (date(2026, 3, 2), "PAYMENT ACME SUPPLIES LTD CHQ-1041", D("1250.00"), None),
    (date(2026, 3, 3), "TRANSFER GLOBEX INDUSTRIAL LLC TRF-2203", D("3400.00"), None),
    (date(2026, 3, 4), "DEPOSIT WAYNE ENTERPRISES RCT-3301", None, D("8000.00")),
    (date(2026, 3, 5), "PAYMENT INITECH OFFICE CO CHQ-1042", D("780.50"), None),
    (date(2026, 3, 6), "DIRECT DEBIT UTILITY BOARD DD-7781", D("412.30"), None),
    (date(2026, 3, 9), "TRANSFER PAYROLL MARCH PAY-0326", D("21500.00"), None),
    (date(2026, 3, 10), "PAYMENT UMBRELLA TRADING FZE CHQ-1043", D("4050.00"), None),
    (date(2026, 3, 11), "TRANSFER HOOLI SERVICES TRF-2204", D("2300.00"), None),
    (date(2026, 3, 12), "PAYMENT STARK INDUSTRIES CHQ-1044", D("1980.00"), None),
    (date(2026, 3, 15), "BANK CHARGES MARCH", D("25.00"), None),
    (date(2026, 3, 17), "DEPOSIT BATCH 0317", None, D("4500.00")),
    (date(2026, 3, 18), "PAYMENT WAYNE LOGISTICS CHQ-1045", D("960.00"), None),
    (date(2026, 3, 20), "TRANSFER ACME SUPPLIES LTD TRF-2205", D("2100.00"), None),
    (date(2026, 3, 24), "DEPOSIT OSCORP RCT-3320", None, D("5250.00")),
    (date(2026, 3, 26), "PAYMENT GLOBEX INDUSTRIAL LLC CHQ-1046", D("1500.00"), None),
    (date(2026, 3, 30), "INTEREST CREDIT", None, D("12.34")),
    (date(2026, 4, 2), "DEPOSIT STARK INDUSTRIES RCT-3330", None, D("3000.00")),
]

# cash book: (type, date, counterparty, reference, amount, expected findings)
CASH_BOOK = [
    ("bank_payment", date(2026, 3, 1), "Acme Supplies Ltd", "CHQ-1041", D("1250.00"), ()),
    ("bank_payment", date(2026, 3, 2), "Globex Industrial LLC", "TRF-2203", D("3400.00"), ()),
    ("bank_receipt", date(2026, 3, 4), "Wayne Enterprises", "RCT-3301", D("8000.00"), ()),
    ("bank_payment", date(2026, 3, 4), "Initech Office Co", "CHQ-1042", D("780.50"), ()),
    ("bank_payment", date(2026, 3, 6), "Utility Board", "DD-7781", D("412.30"), ()),
    ("bank_payment", date(2026, 3, 9), "Payroll March", "PAY-0326", D("21500.00"), ()),
    ("bank_payment", date(2026, 3, 9), "Umbrella Trading FZE", "CHQ-1043", D("4500.00"), ("AMOUNT_MISMATCH",)),      # statement: 4,050.00
    ("bank_payment", date(2026, 3, 10), "Initech Office Co", "TRF-2204", D("2300.00"), ("COUNTERPARTY_MISMATCH",)),  # statement: Hooli Services
    ("bank_payment", date(2026, 3, 11), "Stark Industries", "CHQ-1044", D("1980.00"), ()),
    ("bank_payment", date(2026, 3, 13), "Stark Industries", "CHQ-1044", D("1980.00"), ("VOUCHER_REUSED",)),          # entered twice
    ("bank_receipt", date(2026, 3, 16), "Cyberdyne Systems", "RCT-3310", D("1500.00"), ()),                         # three receipts, one
    ("bank_receipt", date(2026, 3, 16), "Soylent Corp", "RCT-3311", D("1800.00"), ()),                              # batched bank deposit
    ("bank_receipt", date(2026, 3, 17), "Tyrell Corp", "RCT-3312", D("1200.00"), ()),                               # of 4,500.00
    ("bank_payment", date(2026, 3, 17), "Wayne Logistics", "CHQ-1045", D("960.00"), ()),
    ("bank_payment", date(2026, 3, 19), "Acme Supplies Ltd", "TRF-2205", D("2100.00"), ()),
    ("bank_receipt", date(2026, 3, 24), "Oscorp", "RCT-3320", D("5250.00"), ()),
    ("bank_payment", date(2026, 3, 22), "Globex Industrial LLC", "CHQ-1046", D("1500.00"), ()),
    ("bank_payment", date(2026, 3, 28), "Tyrell Corp", "CHQ-1047", D("1100.00"), ("NO_VOUCHER",)),                   # cheque not yet presented
    ("bank_receipt", date(2026, 3, 31), "Stark Industries", "RCT-3330", D("3000.00"), ("CUTOFF_ISSUE",)),            # bank credited 2 April
]


def _rows():
    bal, out = OPENING, []
    for d, desc, dr, cr in STATEMENT:
        bal += (cr or D(0)) - (dr or D(0))
        out.append((d, desc, dr, cr, bal))
    return out


def _pdf_lines(rows) -> list[str]:
    lines = ["FIRST DEMO BANK PJSC", "ACCOUNT STATEMENT", "Customer: Contoso Ltd", f"Account number: {ACCOUNT}", "Currency: USD",
             "Statement period: 01 Mar 2026 to 15 Mar 2026", f"Opening balance 01 Mar 2026     {OPENING:,.2f}", "",
             "Date          Description                                          Debit            Credit            Balance"]
    for d, desc, dr, cr, bal in rows:
        lines.append(f"{d.day:02d} {d.strftime('%b %Y')}   {desc:<44}   {(f'{dr:,.2f}' if dr else ''):>12}   {(f'{cr:,.2f}' if cr else ''):>12}   {bal:>12,.2f}")
    lines += ["", f"Closing balance 15 Mar 2026     {rows[-1][4]:,.2f}" if rows else "", "End of statement part 1"]
    return lines


def generate_bank_demo(out: Path) -> tuple[Path, Path, dict[str, set[str]]]:
    import openpyxl
    from openpyxl.styles import Font

    ev = out / "evidence"
    ev.mkdir(parents=True, exist_ok=True)
    rows = _rows()
    part1 = [r for r in rows if r[0] <= date(2026, 3, 15)]
    part2 = [r for r in rows if r[0] > date(2026, 3, 15)]
    (ev / "statement_2026-03-01_to_03-15.pdf").write_bytes(make_pdf(_pdf_lines(part1)))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["FIRST DEMO BANK PJSC - ACCOUNT STATEMENT"])
    ws.append([f"Account number: {ACCOUNT}", "Currency: USD", "Period: 16 Mar 2026 to 07 Apr 2026"])
    ws.append([])
    ws.append(["Date", "Description", "Debit", "Credit", "Balance"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for d, desc, dr, cr, bal in part2:
        ws.append([d, desc, float(dr) if dr else None, float(cr) if cr else None, float(bal)])
    ws.column_dimensions["B"].width = 46
    wb.save(ev / "statement_2026-03-16_to_04-07.xlsx")

    gl = openpyxl.Workbook()
    g = gl.active
    g.title = "Cash book"
    g.append(["Type", "Date", "Counterparty", "Reference", "Amount", "Currency", "Description"])
    expected: dict[str, set[str]] = {}
    for i, (typ, d, party, ref, amt, codes) in enumerate(CASH_BOOK):
        g.append([typ, d, party, ref, float(amt), "USD", f"{'Payment to' if typ == 'bank_payment' else 'Receipt from'} {party}"])
        if codes:
            expected[f"Transaction:R{i + 2}"] = set(codes)
    for c in g[1]:
        c.font = Font(bold=True)
    ledger = out / "cash_book_march_2026.xlsx"
    gl.save(ledger)

    lines = ["item,expected"] + [f"R{i + 2},{';'.join(sorted(expected.get(f'Transaction:R{i + 2}', ()))) or 'CLEAN'}" for i in range(len(CASH_BOOK))]
    (out / "answer_key.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return ledger, ev, expected

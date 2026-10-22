"""Generate a realistic mixed-format demo: ledger (xlsx) + supporting documents (txt/docx/xlsx/pdf/eml/png).

Seeded errors give a golden answer key: missing document, amount mismatch, wrong counterparty,
cut-off, reused document, missing approval, an unreadable scan, and an orphan document.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

PERIOD = (date(2026, 1, 1), date(2026, 3, 31))


def make_pdf(lines: list[str], per_page: int = 46) -> bytes:
    """Minimal text PDF (no external library); long listings continue on further pages."""
    esc = lambda x: x.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    pages = [lines[i:i + per_page] for i in range(0, max(len(lines), 1), per_page)]
    n = len(pages)
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            "<< /Type /Pages /Kids [" + " ".join(f"{3 + 2 * i} 0 R" for i in range(n)) + f"] /Count {n} >>"]
    font_id = 3 + 2 * n
    for i, chunk in enumerate(pages):
        content = "BT /F1 10 Tf 40 800 Td 15 TL\n" + "".join(f"({esc(l)}) Tj T*\n" for l in chunk) + "ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>")
        objs.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def _invoice(vendor, number, dt, bill_to, total, *, tax_rate="0.05", approver=None, title="TAX INVOICE",
             party_label="Supplier") -> list[str]:
    from decimal import Decimal

    total = Decimal(total)
    sub = (total / (1 + Decimal(tax_rate))).quantize(Decimal("0.01"))
    tax = total - sub
    lines = [vendor, title, f"Invoice No: {number}", f"Invoice Date: {dt}", "Due Date: 30 days from invoice",
             f"Bill To: {bill_to}", f"{party_label}: {vendor}", "Description\tQty\tAmount",
             f"Goods and services\t1\t{sub:,.2f}", f"Subtotal\t\t{sub:,.2f}", f"VAT {int(Decimal(tax_rate) * 100)}%\t\t{tax:,.2f}",
             f"Total Amount Due: USD {total:,.2f}"]
    if approver:
        lines.append(f"Approved by: {approver}")
    return lines


# (type, date, counterparty, reference, amount, approver, doc format, expected codes)
LEDGER = [
    ("purchase", date(2026, 2, 3), "Acme Supplies Ltd", "INV-5001", "1250.00", "", "txt", ()),
    ("purchase", date(2026, 2, 5), "Globex Industrial LLC", "INV-5002", "3400.00", "", "docx", ()),
    ("purchase", date(2026, 2, 9), "Initech Office Co", "INV-5003", "780.50", "", "xlsx", ()),
    ("purchase", date(2026, 2, 12), "Acme Supplies Ltd", "INV-5004", "2100.00", "", "pdf", ()),
    ("sales", date(2026, 2, 15), "Wayne Enterprises", "SI-9001", "8000.00", "", "txt", ()),
    ("expense", date(2026, 2, 16), "Cafe Roma", "EXP-301", "142.75", "", "eml", ()),
    ("payroll", date(2026, 2, 28), "Dana Lee", "PAY-0226", "4250.00", "", "payslip", ()),
    ("fixed_asset", date(2026, 3, 2), "Umbrella Trading FZE", "FA-77", "15000.00", "R. Khan", "txt", ()),
    # --- seeded errors ---
    ("purchase", date(2026, 3, 3), "Acme Supplies Ltd", "INV-5010", "960.00", "", None, ("NO_VOUCHER",)),
    ("purchase", date(2026, 3, 4), "Globex Industrial LLC", "INV-5011", "1500.00", "", "txt:1800.00", ("AMOUNT_MISMATCH",)),
    ("purchase", date(2026, 3, 5), "Initech Office Co", "INV-5012", "600.00", "", "txt:vendor=Hooli Services", ("COUNTERPARTY_MISMATCH",)),
    ("purchase", date(2026, 3, 30), "Acme Supplies Ltd", "INV-5013", "2000.00", "", "txt:date=12 Apr 2026", ("CUTOFF_ISSUE",)),
    ("purchase", date(2026, 3, 9), "Initech Office Co", "INV-5014", "700.00", "", "txt", ()),
    ("purchase", date(2026, 3, 16), "Initech Office Co", "INV-5014", "700.00", "", None, ("VOUCHER_REUSED",)),
    ("fixed_asset", date(2026, 3, 18), "Umbrella Trading FZE", "FA-78", "22000.00", "", "txt", ("MISSING_APPROVAL",)),
    ("purchase", date(2026, 3, 20), "Acme Supplies Ltd", "INV-5016", "410.00", "", "png", ("NO_VOUCHER",)),
]


def generate_demo(out: Path) -> tuple[Path, Path, dict[str, set[str]]]:
    import docx
    import openpyxl

    docs = out / "docs"
    docs.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GL"
    ws.append(["Type", "Date", "Counterparty", "Reference", "Amount", "Approver", "Description"])
    expected: dict[str, set[str]] = {}
    for i, (typ, dt, party, ref, amount, approver, fmt, codes) in enumerate(LEDGER):
        row = i + 2
        ws.append([typ, dt, party, ref, float(amount), approver, f"{typ} {ref}"])
        if codes:
            expected[f"Transaction:R{row}"] = set(codes)
        if fmt is None:
            continue
        kind, _, override = fmt.partition(":")
        shown_total, vendor, shown_date = amount, party, f"{dt.day:02d} {dt.strftime('%b %Y')}"
        if override:
            if override.startswith("vendor="):
                vendor = override.split("=", 1)[1]
            elif override.startswith("date="):
                shown_date = override.split("=", 1)[1]
            else:
                shown_total = override
        stem = ref.lower().replace("-", "_")
        if typ == "sales":
            lines = _invoice("Our Company LLC", ref, shown_date, party, shown_total, party_label="Issued by")
        elif typ == "expense":
            lines = None
        elif typ == "payroll":
            lines = ["Our Company LLC", "PAYSLIP", f"Payslip No: {ref}", f"Pay Date: {shown_date}",
                     f"Employee Name: {party}", "Gross Pay\t4,800.00", "Deductions\t550.00", f"Net Pay: USD {float(shown_total):,.2f}"]
        else:
            lines = _invoice(vendor, ref, shown_date, "Our Company LLC", shown_total,
                             approver=approver or None)
        if kind == "eml":
            (docs / f"{stem}.eml").write_text(
                f"From: receipts@caferoma.example\nTo: claims@ourcompany.example\nDate: Mon, 16 Feb 2026 12:00:00 +0000\n"
                f"Subject: Receipt {ref}\n\nMerchant: {party}\nReceipt No: {ref}\nEmployee: Sam Patel\n"
                f"Total: USD {float(shown_total):,.2f}\n", encoding="utf-8")
        elif kind == "txt":
            (docs / f"{stem}.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
        elif kind == "payslip":
            (docs / f"{stem}.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
        elif kind == "docx":
            d = docx.Document()
            for l in lines:
                d.add_paragraph(l)
            d.save(docs / f"{stem}.docx")
        elif kind == "xlsx":
            w = openpyxl.Workbook()
            for l in lines:
                w.active.append(l.split("\t") if "\t" in l else [l])
            w.save(docs / f"{stem}.xlsx")
        elif kind == "pdf":
            (docs / f"{stem}.pdf").write_bytes(make_pdf([l.replace("\t", "   ") for l in lines]))
        elif kind == "png":  # a "scan": image bytes with no text layer
            (docs / f"{stem}.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    # an invoice with no ledger entry (possible unrecorded liability)
    (docs / "orphan_inv_5099.txt").write_text("\n".join(_invoice("Acme Supplies Ltd", "INV-5099", "20 Mar 2026",
                                                                  "Our Company LLC", "333.00")) + "\n", encoding="utf-8")
    ledger = out / "ledger.xlsx"
    wb.save(ledger)
    return ledger, docs, expected

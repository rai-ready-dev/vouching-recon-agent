"""Workpaper exports: Markdown, Word and Excel, including reviewer decisions and sign-off."""
from __future__ import annotations

import io
import json
from datetime import datetime

from ..evidence import EvidenceRef
from ..reports import RunReport
from ..workpaper import render_workpaper
from .service import Service

STATUS_LABEL = {"resolved": "Resolved", "confirmed": "Confirmed exception", "followup": "Follow-up needed",
                "pending": "Pending"}


def load_bundle(svc: Service, rid: int) -> dict:
    run = svc.run_public(rid)
    if run["status"] != "completed":
        from .service import ApiError

        raise ApiError(409, "The run has not completed.")
    wd = svc.settings.runs_dir / str(rid)
    report = RunReport.model_validate_json((wd / "report.json").read_text())
    manifest = [EvidenceRef(**m) for m in json.loads((wd / "manifest.json").read_text())]
    docs = json.loads((wd / "documents.json").read_text())
    rows = svc.db.query("SELECT txn_id FROM items WHERE run_id=? ORDER BY txn_id", (rid,))
    items = []
    for r in rows:
        d = svc.item_detail(rid, r["txn_id"])
        items.append(d)
    eng = svc.engagement(run["engagement_id"])
    return {"run": run, "eng": eng, "report": report, "manifest": manifest, "docs": docs, "items": items}


def _final_note(item: dict) -> tuple[str, str, str]:
    """(status label, reviewer, note) from the last status-changing decision."""
    last = [d for d in item["decisions"] if d["action"] != "comment"]
    if not last:
        return STATUS_LABEL["pending"], "", ""
    d = last[-1]
    return STATUS_LABEL.get(d["action"], d["action"]), d.get("user") or "", d.get("note") or ""


def _needs_attention(item: dict) -> bool:
    return item["disposition"] in ("review", "escalate")


def build_markdown(svc: Service, rid: int) -> str:
    b = load_bundle(svc, rid)
    run, eng = b["run"], b["eng"]
    head = [f"# {eng['name']}" + (f" - {eng['client']}" if eng["client"] else ""), "",
            f"Run {rid} | started by {run.get('started_by_name', '')} | {run['started_at']}", ""]
    out = "\n".join(head) + render_workpaper(b["report"], manifest=b["manifest"])
    out = out.replace("## Conclusion (draft)", "## Reviewer decisions\n\n" + _decisions_md(b) + "\n## Conclusion (draft)")
    if run["locked"]:
        out += (f"\n## Sign-off\n\nSigned off by **{run.get('signoff_by_name', '')}** on {run['signoff_at']}."
                + (f"\n\n> {run['signoff_note']}" if run["signoff_note"] else "") + "\n")
    else:
        out += "\n## Sign-off\n\n**Not signed off.** This workpaper is a draft.\n"
    return out


def _decisions_md(b: dict) -> str:
    lines = ["| Item | Counterparty | Amount | Findings | Status | Reviewer | Note |", "|---|---|---|---|---|---|---|"]
    for it in b["items"]:
        if not _needs_attention(it):
            continue
        t = it["transaction"]
        status, who, note = _final_note(it)
        codes = ", ".join(sorted({f["code"] for f in it["findings"]})) or "sample check"
        lines.append(f"| {it['txn_id']} | {t['counterparty']} | {float(t['amount']):,.2f} | {codes} | {status} | {who} | "
                     f"{note.replace('|', '/').replace(chr(10), ' ')} |")
    return "\n".join(lines) + "\n" if len(lines) > 2 else "No items required a decision.\n"


def build_docx(svc: Service, rid: int) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    b = load_bundle(svc, rid)
    run, eng, rep = b["run"], b["eng"], b["report"]
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(f"Vouching workpaper: {eng['name']}", 0)
    meta = [("Client", eng["client"] or "-"), ("Audit period", rep.period or "-"),
            ("Prepared by", run.get("started_by_name", "")), ("Run", f"{rid} ({run['started_at']})"),
            ("Reasoning layer", rep.analyst), ("Policy version", rep.policy_version),
            ("Population", f"{rep.population} ledger entries, value {rep.population_value:,.2f}"),
            ("Status", f"Signed off by {run.get('signoff_by_name')} on {run['signoff_at']}" if run["locked"] else "DRAFT - not signed off")]
    t = d.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in meta:
        r = t.add_row().cells
        r[0].text, r[1].text = k, str(v)
    counts = rep.by_disposition()
    d.add_heading("Result summary", 1)
    d.add_paragraph(f"Cleared by tests: {counts['auto_cleared']}   |   For review: {counts['review']}   |   "
                    f"Escalated: {counts['escalate']}   |   Exceptions: {len(rep.cases)}")
    d.add_heading("Procedure performed", 1)
    for line in ("Every supporting document was read and key fields extracted; each extracted value was verified to "
                 "appear in the document.",
                 "Each ledger entry was linked to its supporting document by reference, amount, counterparty and date.",
                 "The vouching programme for each transaction type was applied: " +
                 "; ".join(f"{k.replace('_', ' ')} ({len(v)} tests)" for k, v in sorted(rep.programmes.items())) + ".",
                 "Documents not linked to any ledger entry were listed as possible unrecorded items."):
        d.add_paragraph(line, style="List Number")

    d.add_heading("Exceptions and reviewer decisions", 1)
    att = [it for it in b["items"] if _needs_attention(it)]
    if att:
        tbl = d.add_table(rows=1, cols=6)
        tbl.style = "Light Grid Accent 1"
        for c, h in zip(tbl.rows[0].cells, ("Item", "Counterparty / ref", "Amount", "Findings", "Status", "Reviewer note")):
            c.text = h
        for it in att:
            tx = it["transaction"]
            status, who, note = _final_note(it)
            r = tbl.add_row().cells
            r[0].text = it["txn_id"].replace("Transaction:", "")
            r[1].text = f"{tx['counterparty']} / {tx['reference']}"
            r[2].text = f"{float(tx['amount']):,.2f}"
            r[3].text = "\n".join(f"{f['code']}: {f['message']}" for f in it["findings"]) or "Sampled (no exception)"
            r[4].text = status + (f"\n{who}" if who else "")
            r[5].text = note
    else:
        d.add_paragraph("No exceptions.")

    if rep.intake_issues:
        d.add_heading("Documents that could not be fully read", 1)
        for i in rep.intake_issues:
            d.add_paragraph(i, style="List Bullet")
    if rep.unlinked_documents:
        d.add_heading("Documents not linked to any ledger entry (possible unrecorded items)", 1)
        for i in rep.unlinked_documents:
            d.add_paragraph(i, style="List Bullet")
    corrected = [(doc["filename"], v) for doc in b["docs"]["documents"] for v in doc["vouchers"] if v["corrected"]]
    if corrected:
        d.add_heading("Extraction corrected by a reviewer", 1)
        for name, v in corrected:
            d.add_paragraph(f"{name}: total {v['amount']}, date {v['date']}, parties {', '.join(v['parties'])}",
                            style="List Bullet")

    d.add_heading("Conclusion", 1)
    pending = sum(1 for it in att if _final_note(it)[0] in ("Pending", "Follow-up needed"))
    confirmed = sum(1 for it in att if _final_note(it)[0] == "Confirmed exception")
    d.add_paragraph(f"{len(att)} item(s) required reviewer attention: {confirmed} confirmed as exceptions, "
                    f"{len(att) - confirmed - pending} resolved, {pending} outstanding. " +
                    ("This workpaper has been signed off." if run["locked"] else "This is a draft until signed off."))
    d.add_heading("Sign-off", 1)
    if run["locked"]:
        d.add_paragraph(f"Reviewed and signed off by {run.get('signoff_by_name')} on {run['signoff_at']}.")
        if run["signoff_note"]:
            d.add_paragraph(run["signoff_note"])
    else:
        d.add_paragraph("Reviewed by: ______________________    Date: ______________")

    d.add_heading("Evidence manifest (SHA-256)", 1)
    mt = d.add_table(rows=1, cols=2)
    mt.rows[0].cells[0].text, mt.rows[0].cells[1].text = "Record", "SHA-256 (first 16)"
    for m in b["manifest"][:400]:
        r = mt.add_row().cells
        r[0].text, r[1].text = m.id, m.sha256[:16]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_xlsx(svc: Service, rid: int) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    b = load_bundle(svc, rid)
    run, eng, rep = b["run"], b["eng"], b["report"]
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, headers, rows, widths=None, money_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(48, len(h) + 4)))
        for col in money_cols:
            for cell in ws[get_column_letter(col)][1:]:
                cell.number_format = "#,##0.00"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    ws = wb.active
    ws.title = "Summary"
    counts = rep.by_disposition()
    for k, v in [("Engagement", eng["name"]), ("Client", eng["client"]), ("Audit period", rep.period),
                 ("Run", rid), ("Prepared by", run.get("started_by_name", "")), ("Started", run["started_at"]),
                 ("Reasoning layer", rep.analyst), ("Policy version", rep.policy_version),
                 ("Ledger entries", rep.population), ("Population value", float(rep.population_value)),
                 ("Cleared by tests", counts["auto_cleared"]), ("For review", counts["review"]),
                 ("Escalated", counts["escalate"]), ("Exceptions", len(rep.cases)),
                 ("Documents with problems", len(rep.intake_issues)), ("Unlinked documents", len(rep.unlinked_documents)),
                 ("Sign-off", f"{run.get('signoff_by_name')} on {run['signoff_at']}" if run["locked"] else "Not signed off")]:
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 60
    for c in ws["A"]:
        c.font = Font(bold=True)

    rows = []
    for it in b["items"]:
        tx = it["transaction"]
        status, who, note = _final_note(it)
        for f in it["findings"] or [None]:
            rows.append([it["txn_id"].replace("Transaction:", ""), tx["type"], tx["date"], tx["counterparty"], tx["reference"],
                         float(tx["amount"]), it["disposition"], f["severity"] if f else "", f["code"] if f else "",
                         f["message"] if f else "", float(f["amount_at_risk"]) if f else None,
                         (it["source_file"] or {}).get("filename", ""), status, who, note])
    sheet("Exceptions", ["Item", "Type", "Date", "Counterparty", "Reference", "Ledger amount", "Disposition", "Severity",
                         "Finding", "Detail", "Amount at risk", "Document", "Review status", "Reviewer", "Note"],
          [r for r in rows if r[8]], {"Detail": 60, "Note": 40, "Counterparty": 24, "Document": 26}, money_cols=(6, 11))
    sheet("All items", ["Item", "Type", "Date", "Counterparty", "Reference", "Ledger amount", "Disposition", "Severity",
                        "Finding", "Detail", "Amount at risk", "Document", "Review status", "Reviewer", "Note"], rows,
          {"Detail": 60, "Note": 40, "Counterparty": 24, "Document": 26}, money_cols=(6, 11))
    drows = []
    for doc in b["docs"]["documents"]:
        if not doc["vouchers"]:
            drows.append([doc["filename"], doc["kind"], doc["method"], "", "", None, "", "", "", "; ".join(doc["issues"])])
        for v in doc["vouchers"]:
            drows.append([doc["filename"], doc["kind"], v["method"], v["doc_type"], v["date"] or "",
                          float(v["amount"]) if v["amount"] is not None else None, "; ".join(v["parties"]),
                          "; ".join(v["references"]), ", ".join(t.replace("Transaction:", "") for t in v["linked_txns"]),
                          "; ".join(doc["issues"])])
    sheet("Documents", ["File", "Kind", "Read by", "Doc type", "Date", "Total", "Parties", "Numbers", "Linked to", "Issues"],
          drows, {"File": 30, "Parties": 30, "Issues": 40}, money_cols=(6,))
    sheet("Not linked", ["Document (matched no ledger entry)"], [[x] for x in rep.unlinked_documents], {"Document (matched no ledger entry)": 60})
    sheet("Unread - issues", ["Issue"], [[x] for x in rep.intake_issues], {"Issue": 120})
    dec = svc.db.query("SELECT d.txn_id, d.action, d.note, d.created_at, u.display_name AS who, d.carried_from FROM decisions d"
                       " LEFT JOIN users u ON u.id=d.user_id WHERE d.run_id=? ORDER BY d.id", (rid,))
    sheet("Decision log", ["Item", "Action", "Note", "When (UTC)", "By", "Carried from run"],
          [[r["txn_id"].replace("Transaction:", ""), r["action"], r["note"], r["created_at"], r["who"], r["carried_from"]] for r in dec],
          {"Note": 70, "When (UTC)": 24})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

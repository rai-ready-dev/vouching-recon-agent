"""Universal document ingestion: pdf, scans/images, xlsx, docx, csv, txt, eml -> text.

Nothing is dropped silently: unreadable, encrypted or image-only files raise
`IngestError` or are flagged `needs_ocr`, and the caller reports them as issues.
"""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from email import policy as email_policy
from email.parser import BytesParser
from pathlib import Path
from typing import Protocol

TEXT_EXT = {".txt", ".md", ".log", ".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SUPPORTED_EXT = TEXT_EXT | IMAGE_EXT | {".pdf", ".xlsx", ".xlsm", ".docx", ".eml"}
MIN_CHARS_PER_PAGE = 20
MAX_SHEET_ROWS = 5000


class IngestError(Exception):
    """File could not be read; the message is shown to the reviewer."""


@dataclass
class SourceDocument:
    id: str
    filename: str
    kind: str
    sha256: str
    text: str = ""
    pages: int = 1
    method: str = "text"
    needs_ocr: bool = False
    raw: bytes | None = None
    warnings: list[str] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)  # integrity observations (edited PDF, changed author...)


class OcrEngine(Protocol):
    name: str

    def ocr(self, doc: SourceDocument) -> str: ...


EDITOR_TOOLS = ("photoshop", "illustrator", "gimp", "nitro", "foxit phantom", "phantompdf", "sejda", "ilovepdf",
                "smallpdf", "pdfescape", "pdf-xchange editor", "canva", "pdf editor", "pdffiller", "docfly", "lightpdf")
EDIT_GAP_DAYS = 1


def _days(a, b) -> int | None:
    try:
        return abs((b.replace(tzinfo=None) - a.replace(tzinfo=None)).days)
    except Exception:
        return None


def _flags_from_dates(kind: str, created, modified) -> list[str]:
    gap = _days(created, modified) if created and modified else None
    if gap is not None and gap > EDIT_GAP_DAYS and modified.replace(tzinfo=None) > created.replace(tzinfo=None):
        return [f"{kind} was modified {gap} days after it was created ({created:%Y-%m-%d} to {modified:%Y-%m-%d})"]
    return []


def _pdf_flags(reader) -> list[str]:
    flags: list[str] = []
    try:
        md = reader.metadata
        if not md:
            return flags
        flags += _flags_from_dates("PDF", md.creation_date, md.modification_date)
        for label, val in (("producer", md.producer), ("creator", md.creator)):
            if val and any(t in str(val).lower() for t in EDITOR_TOOLS):
                flags.append(f"PDF {label} is an editing tool ('{val}')")
                break
    except Exception:  # malformed metadata is common and not itself suspicious
        pass
    return flags


def _office_flags(kind: str, created, modified, author, last_by) -> list[str]:
    flags = _flags_from_dates(kind, created, modified)
    if author and last_by and str(author).strip().lower() != str(last_by).strip().lower():
        flags.append(f"{kind} author is '{author}' but it was last saved by '{last_by}'")
    return flags


def _decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int,)) and not isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        d = Decimal(str(v))
        return f"{d.quantize(Decimal('0.01')):f}" if d.as_tuple().exponent >= -2 else f"{d:f}"
    return str(v).strip()


def _xlsx_text(raw: bytes) -> tuple[str, int, list[str]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise IngestError(f"cannot open workbook ({type(exc).__name__}); it may be corrupt or password-protected") from exc
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"[Sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                out.append(f"... truncated after {MAX_SHEET_ROWS} rows")
                break
            cells = [_cell(c) for c in row]
            if any(cells):
                out.append("\t".join(cells).rstrip())
    p = wb.properties
    flags = _office_flags("Workbook", p.created, p.modified, p.creator, p.lastModifiedBy) if p else []
    return "\n".join(out), len(wb.worksheets), flags


def _docx_text(raw: bytes) -> tuple[str, list[str]]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise IngestError(f"cannot open Word file ({type(exc).__name__})") from exc
    lines: list[str] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, d).text)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                lines.append("\t".join(c.text.strip() for c in row.cells))
    cp = d.core_properties
    flags = _office_flags("Word file", cp.created, cp.modified, cp.author, cp.last_modified_by)
    return "\n".join(lines), flags


def _pdf_text(raw: bytes) -> tuple[list[str], list[str], list[str]]:
    from pypdf import PdfReader

    warnings: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestError("PDF is password-protected")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"cannot read PDF ({type(exc).__name__}: {exc})") from exc
    return pages, warnings, _pdf_flags(reader)


def _eml_text(raw: bytes) -> tuple[str, list[str]]:
    msg = BytesParser(policy=email_policy.default).parsebytes(raw)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Date", "Subject") if msg[h])
    n_att = sum(1 for _ in msg.iter_attachments())
    warns = [f"{n_att} email attachment(s) were not read; supply them as separate files"] if n_att else []
    return f"{head}\n\n{text}", warns


def ingest_path(path: str | Path, root: str | Path | None = None) -> SourceDocument:
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXT:
        raise IngestError(f"unsupported file type '{ext or p.name}'")
    try:
        raw = p.read_bytes()
    except OSError as exc:
        raise IngestError(f"cannot read file: {exc}") from exc
    if not raw:
        raise IngestError("file is empty")
    rel = p.relative_to(root).as_posix() if root else p.name
    doc = SourceDocument(id=rel, filename=rel, kind=ext.lstrip("."), sha256=hashlib.sha256(raw).hexdigest())

    if ext in TEXT_EXT:
        doc.text = _decode(raw)
    elif ext in (".xlsx", ".xlsm"):
        doc.text, doc.pages, doc.flags = _xlsx_text(raw)
        doc.kind = "xlsx"
    elif ext == ".docx":
        doc.text, doc.flags = _docx_text(raw)
    elif ext == ".eml":
        doc.text, doc.warnings = _eml_text(raw)
    elif ext == ".pdf":
        pages, doc.warnings, doc.flags = _pdf_text(raw)
        doc.pages = max(1, len(pages))
        doc.text = "\n".join(pages)
        thin = [i + 1 for i, t in enumerate(pages) if len(t.strip()) < MIN_CHARS_PER_PAGE]
        if pages and len(thin) == len(pages):
            doc.needs_ocr, doc.raw = True, raw
        elif thin:
            doc.warnings.append(f"page(s) {thin} have no text layer (scanned?); content may be incomplete")
    else:  # image
        doc.kind, doc.needs_ocr, doc.raw = "image", True, raw
    return doc
